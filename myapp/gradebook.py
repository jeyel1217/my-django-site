"""Gradebook: every student of one classroom against every published quiz.

Rules (the same ones the student's own progress page uses, so the numbers agree):
  * a quiz score is the student's BEST finalized attempt, as a percentage
  * an attempt still waiting for essay grading shows "Pending" and is not counted
  * a quiz past its deadline with no attempt shows "Missing" and is not counted
    in the average (it is counted in the Missing column, so nothing is hidden)
  * an excused student (exemption) shows "Excused"; a quiz not due yet stays blank
  * the Average is the mean of the finalized quiz percentages; the letter grade comes
    from GRADE_SCALE in settings.py (default below) and the scale is printed on the sheet

Only active students of the classroom are listed. Who may download it is decided in the view.
"""
import csv
import io
import statistics

from django.conf import settings
from django.utils import timezone

from . import exams
from . import progress as prog
from .models import Enrollment, Quiz, QuizAccommodation, QuizAttempt

DEFAULT_SCALE = ((90, 'A'), (80, 'B'), (70, 'C'), (60, 'D'), (0, 'F'))

# Cell kinds
SCORE, PENDING, MISSING, EXCUSED, OPEN = 'score', 'pending', 'missing', 'excused', 'open'
LABELS = {PENDING: 'Pending', MISSING: 'Missing', EXCUSED: 'Excused', OPEN: ''}


def grade_scale():
    scale = getattr(settings, 'GRADE_SCALE', None) or DEFAULT_SCALE
    return sorted(((float(m), str(l)) for m, l in scale), reverse=True)


def letter_for(pct):
    if pct is None:
        return ''
    for minimum, letter in grade_scale():
        if pct >= minimum:
            return letter
    return grade_scale()[-1][1]


def scale_text():
    scale = grade_scale()
    parts = []
    for i, (minimum, letter) in enumerate(scale):
        upper = f'{scale[i - 1][0] - 0.01:.2f}'.rstrip('0').rstrip('.') if i else '100'
        parts.append(f'{letter} = {minimum:g}% and above' if i == 0 else f'{letter} = {minimum:g}% to {upper}%')
    return '   |   '.join(parts)


def _stats(values):
    values = [v for v in values if v is not None]
    if not values:
        return {'average': None, 'median': None, 'stdev': None}
    return {
        'average': round(sum(values) / len(values), 1),
        'median': round(statistics.median(values), 1),
        'stdev': round(statistics.stdev(values), 1) if len(values) > 1 else None,
    }


def build(room, now=None):
    """Plain data for the sheet; nothing here depends on the file format."""
    now = now or timezone.now()
    quizzes = list(Quiz.objects.filter(classroom=room, is_published=True).order_by('created_at', 'id'))
    enrollments = list(Enrollment.objects.filter(classroom=room, is_active=True)
                       .select_related('student__profile')
                       .order_by('student__profile__name', 'student__username'))
    students = [e.student for e in enrollments]

    by_pair = {}
    for a in (QuizAttempt.objects.filter(quiz__in=quizzes, user__in=students).order_by('taken_at', 'id')):
        by_pair.setdefault((a.user_id, a.quiz_id), []).append(a)
    accs = {(a.student_id, a.quiz_id): a for a in QuizAccommodation.objects.filter(quiz__in=quizzes, student__in=students)}

    rows = []
    for number, e in enumerate(enrollments, start=1):
        s = e.student
        cells, finals, missing, pending, submitted = [], [], 0, 0, 0
        for q in quizzes:
            mine = by_pair.get((s.id, q.id), [])
            first = mine[0] if mine else None
            state = exams.student_quiz_state(q, first, now, accs.get((s.id, q.id)))
            if mine:
                submitted += 1
                done = [a for a in mine if prog.attempt_grading_status(a) == prog.GRADED]
                if done:
                    best = max(done, key=lambda a: (a.score, -a.id))
                    cells.append({'kind': SCORE, 'value': best.score})
                    finals.append(best.score)
                else:
                    cells.append({'kind': PENDING, 'value': None})
                    pending += 1
            elif state['code'] == 'exempt':
                cells.append({'kind': EXCUSED, 'value': None})
            elif state['code'] == 'overdue':
                cells.append({'kind': MISSING, 'value': None})
                missing += 1
            else:
                cells.append({'kind': OPEN, 'value': None})
        average = round(sum(finals) / len(finals), 1) if finals else None
        rows.append({
            'number': number, 'student_id': s.profile.student_id or '', 'name': s.profile.display_name,
            'cells': cells, 'average': average, 'grade': letter_for(sum(finals) / len(finals)) if finals else '',
            'submitted': submitted, 'missing': missing, 'pending': pending,
        })

    column_stats = [_stats([r['cells'][i]['value'] for r in rows if r['cells'][i]['kind'] == SCORE])
                    for i in range(len(quizzes))]
    return {
        'room': room, 'quizzes': quizzes, 'rows': rows, 'column_stats': column_stats,
        'average_stats': _stats([r['average'] for r in rows]),
        'generated': timezone.localtime(now),
        'mentor': room.mentor_name,
        'school': getattr(settings, 'SCHOOL_NAME', 'COLEGIO DE SAN GABRIEL ARCANGEL'),
    }


# ---------------------------------------------------------------------------
# Safe text: a name or title that starts with = + - @ would run as a formula in Excel.
# ---------------------------------------------------------------------------
def safe(text):
    text = '' if text is None else str(text)
    return "'" + text if text[:1] in ('=', '+', '-', '@', '\t', '\r') else text


def _quiz_header(q):
    due = timezone.localtime(q.deadline_datetime).strftime('%b %d') if q.deadline_datetime else 'No deadline'
    return q.title, due


def _plain_text(cell, value):
    """Put user-written text in an Excel cell as text, never as a formula."""
    cell.value = '' if value is None else str(value)
    cell.data_type = 's'
    return cell


# ---------------------------------------------------------------------------
# CSV (opens in Excel too; no extra package needed)
# ---------------------------------------------------------------------------
def to_csv(data):
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(['Gradebook'])
    w.writerow(['School', safe(data['school'])])
    w.writerow(['Classroom', safe(data['room'].name)])
    w.writerow(['Mentor', safe(data['mentor'])])
    w.writerow(['Generated', data['generated'].strftime('%Y-%m-%d %H:%M') + ' (Philippine time)'])
    w.writerow([])
    w.writerow(['No.', 'Student ID', 'Student'] + [safe(q.title) for q in data['quizzes']]
               + ['Average %', 'Grade', 'Submitted', 'Missing', 'Pending'])
    for r in data['rows']:
        w.writerow([r['number'], safe(r['student_id']), safe(r['name'])]
                   + [c['value'] if c['kind'] == SCORE else LABELS[c['kind']] for c in r['cells']]
                   + ['' if r['average'] is None else r['average'], r['grade'],
                      f"{r['submitted']}/{len(data['quizzes'])}", r['missing'], r['pending']])
    for label, key in (('Class average', 'average'), ('Median', 'median'), ('Std. deviation', 'stdev')):
        w.writerow(['', '', label] + ['' if s[key] is None else s[key] for s in data['column_stats']]
                   + ['' if data['average_stats'][key] is None else data['average_stats'][key]])
    w.writerow([])
    w.writerow(['Grading scale', scale_text()])
    return '﻿' + out.getvalue()          # BOM so Excel reads Filipino names correctly


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def to_xlsx(data):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Gradebook'
    head_fill = PatternFill('solid', fgColor='E3CFC4')
    soft_fill = PatternFill('solid', fgColor='F3E9E4')
    thin = Side(style='thin', color='BFBFBF')
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    quizzes = data['quizzes']
    first_quiz_col = 4
    avg_col = first_quiz_col + len(quizzes)
    last_col = avg_col + 4

    ws['A1'] = 'Gradebook'
    ws['A1'].font = Font(size=20, bold=True, color='5A3A2E')
    _plain_text(ws['A2'], data['school'])
    ws['A2'].font = Font(bold=True)
    _plain_text(ws['A3'], 'Classroom: ' + data['room'].name)
    _plain_text(ws['A4'], 'Mentor: ' + data['mentor'])
    ws['A5'] = 'Generated: ' + data['generated'].strftime('%B %d, %Y %I:%M %p') + ' (Philippine time)'
    for r in (3, 4, 5):
        ws.cell(r, 1).font = Font(color='444444')

    hdr = 7
    labels = ['No.', 'Student ID', 'Student']
    for i, text in enumerate(labels, start=1):
        c = ws.cell(hdr, i, text)
        c.font, c.fill, c.alignment, c.border = Font(bold=True), head_fill, center, box
        ws.merge_cells(start_row=hdr, start_column=i, end_row=hdr + 1, end_column=i)
    for j, q in enumerate(quizzes):
        title, due = _quiz_header(q)
        top = _plain_text(ws.cell(hdr, first_quiz_col + j), title)
        top.font, top.fill, top.alignment, top.border = Font(bold=True), head_fill, center, box
        sub = ws.cell(hdr + 1, first_quiz_col + j, 'Due: ' + due if q.deadline_datetime else due)
        sub.font, sub.fill, sub.alignment, sub.border = Font(size=8, color='555555'), head_fill, center, box
    for k, text in enumerate(['Average %', 'Grade', 'Submitted', 'Missing', 'Pending']):
        col = avg_col + k
        c = ws.cell(hdr, col, text)
        c.font, c.fill, c.alignment, c.border = Font(bold=True), head_fill, center, box
        ws.merge_cells(start_row=hdr, start_column=col, end_row=hdr + 1, end_column=col)
    ws.row_dimensions[hdr].height = 42
    ws.row_dimensions[hdr + 1].height = 16

    r = hdr + 2
    n_quiz = len(quizzes)
    for row in data['rows']:
        vals = [row['number'], row['student_id'], row['name']]
        for col, v in enumerate(vals, start=1):
            c = ws.cell(r, col)
            if col == 1:
                c.value = v
            else:
                _plain_text(c, v)
            c.border = box
            c.alignment = Alignment(horizontal='left' if col == 3 else 'center', vertical='center')
        for j, cell in enumerate(row['cells']):
            c = ws.cell(r, first_quiz_col + j)
            c.border, c.alignment = box, Alignment(horizontal='center', vertical='center')
            if cell['kind'] == SCORE:
                c.value = cell['value']
                c.number_format = '0"%"'
                if cell['value'] < 60:
                    c.font = Font(color='B91C1C')
            elif cell['kind'] != OPEN:
                c.value = LABELS[cell['kind']]
                c.font = Font(italic=True, color={'missing': 'B91C1C', 'pending': 'B45309', 'excused': '6D28D9'}[cell['kind']])
        c = ws.cell(r, avg_col, row['average'])
        c.number_format = '0.0"%"'
        ws.cell(r, avg_col + 1, row['grade']).font = Font(bold=True)
        ws.cell(r, avg_col + 2, f"{row['submitted']}/{n_quiz}")
        ws.cell(r, avg_col + 3, row['missing'])
        ws.cell(r, avg_col + 4, row['pending'])
        for col in range(avg_col, last_col + 1):
            ws.cell(r, col).border = box
            ws.cell(r, col).alignment = Alignment(horizontal='center', vertical='center')
            ws.cell(r, col).fill = soft_fill
        r += 1
    if not data['rows']:
        ws.cell(r, 3, 'No students have joined this classroom yet.').font = Font(italic=True, color='777777')
        r += 1

    r += 1
    for label, key in (('Class average', 'average'), ('Median', 'median'), ('Std. deviation', 'stdev')):
        lab = ws.cell(r, 3, label)
        lab.font, lab.alignment = Font(bold=True), Alignment(horizontal='right')
        for j, s in enumerate(data['column_stats']):
            c = ws.cell(r, first_quiz_col + j, s[key])
            c.number_format = '0.0"%"'
            c.alignment, c.border = Alignment(horizontal='center'), box
            c.fill = soft_fill
        c = ws.cell(r, avg_col, data['average_stats'][key])
        c.number_format = '0.0"%"'
        c.alignment, c.border, c.fill = Alignment(horizontal='center'), box, soft_fill
        r += 1

    r += 1
    notes = [
        'Grading scale: ' + scale_text(),
        'The score of a quiz is the best finalized attempt. Average % is the mean of the finalized quiz scores.',
        'Pending = an essay is waiting for the mentor to grade it (not counted yet). Missing = past the deadline with no submission (not counted in the average).',
        'Excused = the mentor excused this student from the quiz. A blank cell = not due yet.',
    ]
    for line in notes:
        ws.cell(r, 1, line).font = Font(size=9, color='555555')
        r += 1

    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 14
    ws.column_dimensions['C'].width = 30
    for j in range(n_quiz):
        ws.column_dimensions[get_column_letter(first_quiz_col + j)].width = 15
    for k, w in enumerate([12, 8, 11, 9, 9]):
        ws.column_dimensions[get_column_letter(avg_col + k)].width = w
    ws.freeze_panes = ws.cell(hdr + 2, first_quiz_col)
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f'{hdr}:{hdr + 1}'

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
