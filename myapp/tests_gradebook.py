import io
from datetime import timedelta

from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from . import gradebook
from .models import QuizAccommodation, QuizAttempt, AuditLog, Question, Choice, Quiz
from .tests_phase2 import Base, FAST, make, quiz


def attempt(user, qz, score, status='finalized'):
    return QuizAttempt.objects.create(user=user, quiz=qz, score=score, total_questions=1, correct_count=1,
                                      grading_status=status)


def sheet(response):
    wb = load_workbook(io.BytesIO(response.content))
    return wb['Gradebook']


def find_row(ws, text):
    for row in ws.iter_rows():
        for c in row:
            if c.value == text:
                return c.row
    raise AssertionError(f'{text!r} not on the sheet')


@FAST
class GradebookTests(Base):
    def setUp(self):
        super().setUp()
        self.s3 = make('student_ccc3', 'STUDENT')
        for s in (self.s1, self.s2, self.s3):
            self.join(s, self.room1)
        self.qa = self.q1                                  # 'Java Quiz'
        self.qb = quiz(self.room1, 'Loops Quiz')
        self.s1.profile.name = 'Ana Cruz'; self.s1.profile.student_id = '2024-001'; self.s1.profile.save()
        self.s2.profile.name = 'Ben Reyes'; self.s2.profile.save()
        self.s3.profile.name = 'Cara Diaz'; self.s3.profile.save()
        attempt(self.s1, self.qa, 80); attempt(self.s1, self.qa, 95)        # best counts
        attempt(self.s1, self.qb, 100)
        attempt(self.s2, self.qa, 50)
        attempt(self.s2, self.qb, 0, status='pending')                      # essay waiting
        # s3: nothing; qa is overdue
        self.qa.deadline_datetime = timezone.now() - timedelta(days=1); self.qa.save()
        self.url = reverse('classroom_gradebook', args=[self.room1.pk])

    def get(self, user, **kw):
        self.login(user)
        return self.client.get(self.url, kw)

    def test_mentor_gets_an_excel_file(self):
        r = self.get(self.m1)
        self.assertEqual(r.status_code, 200)
        self.assertIn('spreadsheetml.sheet', r['Content-Type'])
        self.assertIn('.xlsx', r['Content-Disposition'])
        self.assertEqual(r['Cache-Control'], 'no-store')
        self.assertEqual(sheet(r)['A1'].value, 'Gradebook')

    def test_numbers_follow_the_progress_rules(self):
        ws = sheet(self.get(self.m1))
        hdr = find_row(ws, 'Student')
        titles = [ws.cell(hdr, c).value for c in range(4, 6)]
        self.assertEqual(titles, ['Java Quiz', 'Loops Quiz'])
        ana, ben, cara = (find_row(ws, n) for n in ('Ana Cruz', 'Ben Reyes', 'Cara Diaz'))
        self.assertEqual([ws.cell(ana, 4).value, ws.cell(ana, 5).value], [95, 100])      # best attempt
        self.assertEqual(ws.cell(ana, 6).value, 97.5)                                    # average
        self.assertEqual(ws.cell(ana, 7).value, 'A')
        self.assertEqual(ws.cell(ben, 4).value, 50)
        self.assertEqual(ws.cell(ben, 5).value, 'Pending')
        self.assertEqual(ws.cell(ben, 6).value, 50)                                      # pending not counted
        self.assertEqual(ws.cell(ben, 7).value, 'F')
        self.assertEqual(ws.cell(cara, 4).value, 'Missing')                              # past deadline
        self.assertIsNone(ws.cell(cara, 5).value)                                        # not due: blank
        self.assertIsNone(ws.cell(cara, 6).value)                                        # no average, not 0
        self.assertEqual(ws.cell(cara, 9).value, 1)                                      # Missing count

    def test_class_statistics(self):
        ws = sheet(self.get(self.m1))
        row = find_row(ws, 'Class average')
        self.assertEqual(ws.cell(row, 4).value, 72.5)            # (95+50)/2
        self.assertEqual(ws.cell(find_row(ws, 'Median'), 4).value, 72.5)
        self.assertEqual(ws.cell(find_row(ws, 'Std. deviation'), 4).value, 31.8)

    def test_excused_student(self):
        QuizAccommodation.objects.create(quiz=self.qa, student=self.s3, kind='EXEMPTION', granted_by=self.m1)
        ws = sheet(self.get(self.m1))
        self.assertEqual(ws.cell(find_row(ws, 'Cara Diaz'), 4).value, 'Excused')

    def test_only_active_students_of_this_classroom(self):
        self.join(self.s2, self.room2)
        make('student_ddd4', 'STUDENT')
        e = self.room1.enrollments.get(student=self.s3); e.is_active = False; e.save()
        ws = sheet(self.get(self.m1))
        names = {c.value for row in ws.iter_rows() for c in row}
        self.assertIn('Ana Cruz', names)
        self.assertNotIn('Cara Diaz', names)

    def test_unpublished_quiz_not_listed(self):
        quiz(self.room1, 'Hidden Draft', published=False)
        ws = sheet(self.get(self.m1))
        self.assertNotIn('Hidden Draft', {c.value for row in ws.iter_rows() for c in row})

    def test_permissions(self):
        self.login(self.m2)                                              # another mentor
        self.assertEqual(self.client.get(self.url).status_code, 404)
        self.login(self.s1)
        self.assertEqual(self.client.get(self.url).status_code, 403)
        self.client.logout()
        self.assertEqual(self.client.get(self.url).status_code, 302)     # to login
        self.assertEqual(self.get(self.admin).status_code, 200)

    def test_download_is_recorded(self):
        self.get(self.m1)
        self.assertTrue(AuditLog.objects.filter(action='gradebook_downloaded', actor=self.m1).exists())

    def test_csv_version(self):
        r = self.get(self.m1, format='csv')
        self.assertIn('text/csv', r['Content-Type'])
        text = r.content.decode('utf-8-sig')
        self.assertIn('Ana Cruz', text)
        self.assertIn('Pending', text)
        self.assertIn('Grading scale', text)

    def test_formula_like_names_stay_text(self):
        self.s3.profile.name = '=HYPERLINK("http://x.test","hi")'; self.s3.profile.save()
        ws = sheet(self.get(self.m1))
        cell = next(c for row in ws.iter_rows() for c in row if c.value and 'HYPERLINK' in str(c.value))
        self.assertEqual(cell.data_type, 's')
        csv_text = self.get(self.m1, format='csv').content.decode('utf-8-sig')
        self.assertIn("'=HYPERLINK", csv_text)

    def test_empty_classroom_still_downloads(self):
        room = type(self.room1).objects.create(name='Empty', mentor=self.m1)
        self.login(self.m1)
        r = self.client.get(reverse('classroom_gradebook', args=[room.pk]))
        self.assertEqual(r.status_code, 200)
        self.assertIn('No students', ' '.join(str(c.value) for row in sheet(r).iter_rows() for c in row if c.value))

    def test_letter_scale_and_setting(self):
        self.assertEqual([gradebook.letter_for(x) for x in (100, 90, 89.9, 80, 70, 60, 59.9, 0)],
                         ['A', 'A', 'B', 'B', 'C', 'D', 'F', 'F'])
        from django.test import override_settings
        with override_settings(GRADE_SCALE=((75, 'Passed'), (0, 'Failed'))):
            self.assertEqual(gradebook.letter_for(75), 'Passed')
            self.assertEqual(gradebook.letter_for(74.9), 'Failed')

    def test_buttons_on_the_classroom_page(self):
        self.login(self.m1)
        r = self.client.get(reverse('mentor_classroom_detail', args=[self.room1.pk]))
        self.assertContains(r, self.url)
        self.login(self.admin)
        r = self.client.get(reverse('admin_classroom_detail', args=[self.room1.pk]))
        self.assertContains(r, self.url)
