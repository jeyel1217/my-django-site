import io
import json
import os
import shutil
import tempfile
import zipfile
from unittest import mock, skipUnless

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from django.urls import reverse

from . import lessonfiles
from .models import Lesson, LessonFile, LessonProgress, Quiz, QuizAttempt, Question, Choice, UserBadge
from .tests_phase2 import Base, FAST, make, lesson


def pdf_bytes(pages=3, password=None):
    from pypdf import PdfWriter
    w = PdfWriter()
    for _ in range(pages):
        w.add_blank_page(width=300, height=200)
    if password:
        w.encrypt(password)
    buf = io.BytesIO()
    w.write(buf)
    return buf.getvalue()


def office_zip(kind='docx', extra=None, drop=None):
    """A minimal zip that looks like a Word / PowerPoint file (enough for the content check)."""
    parts = {'[Content_Types].xml': '<Types/>', 'word/document.xml' if kind == 'docx' else 'ppt/presentation.xml': '<x/>'}
    parts.update(extra or {})
    for d in (drop or []):
        parts.pop(d, None)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w') as z:
        for k, v in parts.items():
            z.writestr(k, v)
    return buf.getvalue()


def upload(name, data, ctype='application/octet-stream'):
    return SimpleUploadedFile(name, data, content_type=ctype)


def soffice_available():
    return lessonfiles.find_converter() is not None


@FAST
class P4CBase(Base):
    def setUp(self):
        super().setUp()
        self._tmp = tempfile.mkdtemp(prefix='private-test-')
        self._ov = override_settings(PRIVATE_MEDIA_ROOT=self._tmp)
        self._ov.enable()
        self.addCleanup(self._cleanup)
        self.join(self.s1, self.room1)
        self.join(self.s2, self.room2)
        self.les = self.l1                       # 'Java Basics', published, in room1 (mentor m1)
        self.lf = lessonfiles.attach(self.les, upload('notes.pdf', pdf_bytes(4)), self.m1)

    def _cleanup(self):
        self._ov.disable()
        shutil.rmtree(self._tmp, ignore_errors=True)

    def api(self, slug, page):
        return self.client.post(reverse('lesson_progress', args=[slug]), json.dumps({'page': page}), content_type='application/json')


class UploadValidationTests(P4CBase):
    def reject(self, name, data, fragment):
        with self.assertRaises(lessonfiles.UploadRejected) as cm:
            lessonfiles.inspect_upload(upload(name, data))
        self.assertIn(fragment, str(cm.exception))

    def test_valid_files_pass(self):
        self.assertEqual(lessonfiles.inspect_upload(upload('a.pdf', pdf_bytes(2))), ('PDF', 2))
        self.assertEqual(lessonfiles.inspect_upload(upload('a.DOCX', office_zip('docx')))[0], 'DOCX')
        self.assertEqual(lessonfiles.inspect_upload(upload('a.pptx', office_zip('pptx')))[0], 'PPTX')

    def test_wrong_extension(self):
        for name in ('x.exe', 'x.html', 'x.php', 'x.txt', 'x.zip', 'x.pdf.exe', 'x'):
            self.reject(name, b'hello', 'Only PDF')

    def test_old_office_formats_explained(self):
        self.reject('x.doc', b'hello', '.doc')
        self.reject('x.ppt', b'hello', '.ppt')

    def test_content_must_match_name(self):
        self.reject('fake.pdf', b'<html>not a pdf</html>', 'not a real PDF')
        self.reject('fake.docx', b'%PDF-1.4 something', 'not a real Word')
        self.reject('fake.pptx', b'PK\x03\x04garbage-not-a-zip', 'not a valid')

    def test_word_file_must_contain_word_parts(self):
        self.reject('x.docx', office_zip('pptx'), 'does not match a Word')
        self.reject('x.pptx', office_zip('docx'), 'does not match a PowerPoint')

    def test_macros_refused(self):
        self.reject('x.docx', office_zip('docx', extra={'word/vbaProject.bin': 'x'}), 'macros')

    def test_unsafe_zip_paths_refused(self):
        self.reject('x.docx', office_zip('docx', extra={'../evil.txt': 'x'}), 'unsafe')

    def test_empty_and_oversize(self):
        self.reject('x.pdf', b'', 'empty')
        with override_settings(LESSON_FILE_MAX_MB=1):
            self.reject('big.pdf', pdf_bytes(1) + b'0' * (1024 * 1024 + 5), 'limit is 1 MB')

    def test_encrypted_and_broken_pdf(self):
        self.reject('x.pdf', pdf_bytes(2, password='secret'), 'password')
        self.reject('x.pdf', b'%PDF-1.4\nthis is cut off', 'not a readable PDF')

    def test_too_many_pages(self):
        with mock.patch.object(lessonfiles, 'MAX_PAGES', 3):
            self.reject('x.pdf', pdf_bytes(4), 'limit is 3')

    def test_display_name_is_cleaned(self):
        self.assertEqual(lessonfiles.clean_display_name('../../etc/passwd.pdf'), 'passwd.pdf')
        self.assertEqual(lessonfiles.clean_display_name('C:\\Users\\x\\<evil>.pdf'), 'evil.pdf')
        self.assertEqual(lessonfiles.clean_display_name('a\x00b\n.pdf'), 'ab.pdf')

    def test_stored_name_is_random_and_private(self):
        lf = lessonfiles.attach(self.les, upload('../../../evil name.pdf', pdf_bytes(2)), self.m1)
        stored = lf.original.name
        self.assertRegex(stored, r'^lessons/[0-9a-f]{32}\.pdf$')
        self.assertNotIn('evil', stored)
        self.assertTrue(os.path.exists(os.path.join(self._tmp, stored)))
        self.assertEqual(lf.original_name, 'evil name.pdf')
        with self.assertRaises(ValueError):
            lf.original.url                                # no public URL exists

    def test_private_folder_is_not_static_or_media(self):
        from django.conf import settings
        root = os.path.abspath(self._tmp)
        for public in (settings.STATIC_ROOT, settings.MEDIA_ROOT):
            self.assertFalse(root.startswith(os.path.abspath(str(public))))

    def test_shipped_default_folder_is_separate_from_static_and_media(self):
        from django.conf import settings
        with override_settings():
            del settings.PRIVATE_MEDIA_ROOT
            from .storage import private_root
            self.assertEqual(private_root().name, 'private_media')


class MentorFormTests(P4CBase):
    def post_lesson(self, url, data, file=None):
        self.login(self.m1)
        payload = {'classroom': self.room1.pk, 'title': 'With file', 'order': 5, 'content': 'text', 'is_published': 'on'}
        payload.update(data)
        if file:
            payload['file'] = file
        return self.client.post(url, payload)

    def test_create_lesson_with_pdf(self):
        r = self.post_lesson(reverse('mentor_lesson_create'), {}, upload('new.pdf', pdf_bytes(5)))
        self.assertEqual(r.status_code, 302)
        les = Lesson.objects.get(title='With file')
        self.assertEqual((les.file.kind, les.file.page_count, les.file.conversion_status), ('PDF', 5, 'READY'))

    def test_bad_upload_is_refused_and_lesson_not_created(self):
        r = self.post_lesson(reverse('mentor_lesson_create'), {}, upload('evil.html', b'<script>1</script>'))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'Only PDF')
        self.assertFalse(Lesson.objects.filter(title='With file').exists())

    def test_replace_and_remove_file_deletes_old_stored_file(self):
        old = self.lf.original.name
        self.post_lesson(reverse('mentor_lesson_edit', args=[self.les.pk]), {'title': 'Java Basics'}, upload('v2.pdf', pdf_bytes(7)))
        lf = LessonFile.objects.get(lesson=self.les)
        self.assertEqual(lf.page_count, 7)
        self.assertFalse(os.path.exists(os.path.join(self._tmp, old)))
        self.assertTrue(os.path.exists(os.path.join(self._tmp, lf.original.name)))
        stored = lf.original.name
        self.post_lesson(reverse('mentor_lesson_edit', args=[self.les.pk]), {'title': 'Java Basics', 'remove_file': 'on'})
        self.assertFalse(LessonFile.objects.filter(lesson=self.les).exists())
        self.assertFalse(os.path.exists(os.path.join(self._tmp, stored)))

    def test_deleting_lesson_deletes_files(self):
        stored = os.path.join(self._tmp, self.lf.original.name)
        self.assertTrue(os.path.exists(stored))
        self.login(self.m1)
        self.client.post(reverse('mentor_lesson_delete', args=[self.les.pk]))
        self.assertFalse(os.path.exists(stored))

    def test_other_mentor_cannot_touch_file_actions(self):
        self.login(self.m2)
        for name in ('mentor_lesson_file_download', 'mentor_lesson_file_retry', 'mentor_lesson_file_remove'):
            r = self.client.post(reverse(name, args=[self.les.pk]))
            self.assertEqual(r.status_code, 404, name)
        self.assertEqual(self.client.get(reverse('mentor_lesson_file_download', args=[self.les.pk])).status_code, 404)
        self.assertTrue(LessonFile.objects.filter(lesson=self.les).exists())

    def test_students_cannot_use_mentor_file_actions(self):
        self.login(self.s1)
        for name in ('mentor_lesson_file_download', 'mentor_lesson_file_remove'):
            self.assertNotEqual(self.client.post(reverse(name, args=[self.les.pk])).status_code, 200)
        self.assertTrue(LessonFile.objects.filter(lesson=self.les).exists())


class AccessControlTests(P4CBase):
    def stream(self, user=None, slug=None):
        if user:
            self.login(user)
        return self.client.get(reverse('lesson_file', args=[slug or self.les.slug]))

    def test_enrolled_student_gets_pdf(self):
        r = self.stream(self.s1)
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r['Content-Type'], 'application/pdf')
        self.assertTrue(b''.join(r.streaming_content).startswith(b'%PDF'))
        self.assertIn('no-store', r['Cache-Control'])
        self.assertEqual(r['X-Content-Type-Options'], 'nosniff')

    def test_anonymous_is_sent_to_login(self):
        r = self.client.get(reverse('lesson_file', args=[self.les.slug]))
        self.assertEqual(r.status_code, 302)
        self.assertIn('login', r['Location'])

    def test_student_of_another_classroom_gets_404(self):
        self.assertEqual(self.stream(self.s2).status_code, 404)

    def test_removed_student_loses_access(self):
        from .models import Enrollment
        Enrollment.objects.filter(student=self.s1).update(is_active=False)
        self.assertEqual(self.stream(self.s1).status_code, 404)

    def test_unpublished_lesson_hidden_from_students_but_not_owner(self):
        Lesson.objects.filter(pk=self.les.pk).update(is_published=False)
        self.assertEqual(self.stream(self.s1).status_code, 404)
        self.assertEqual(self.stream(self.m1).status_code, 200)

    def test_other_mentor_gets_404_and_admin_ok(self):
        self.assertEqual(self.stream(self.m2).status_code, 404)
        self.assertEqual(self.stream(self.admin).status_code, 200)

    def test_original_download_only_for_owner_and_admin(self):
        self.login(self.m1)
        r = self.client.get(reverse('mentor_lesson_file_download', args=[self.les.pk]))
        self.assertEqual(r.status_code, 200)
        self.assertIn('attachment', r['Content-Disposition'])
        self.login(self.admin)
        self.assertEqual(self.client.get(reverse('mentor_lesson_file_download', args=[self.les.pk])).status_code, 200)

    def test_file_has_no_public_url_anywhere_on_the_page(self):
        self.login(self.s1)
        html = self.client.get(reverse('lesson_detail', args=[self.les.slug])).content.decode()
        self.assertNotIn(self.lf.original.name, html)
        self.assertNotIn('private_media', html)

    def test_lesson_detail_hidden_from_wrong_student(self):
        self.login(self.s2)
        self.assertEqual(self.client.get(reverse('lesson_detail', args=[self.les.slug])).status_code, 404)

    def test_reading_pdf_of_a_non_viewable_file_is_404(self):
        LessonFile.objects.filter(pk=self.lf.pk).update(conversion_status='UNAVAILABLE', page_count=0)
        self.assertEqual(self.stream(self.s1).status_code, 404)


class ProgressTests(P4CBase):
    def test_pages_must_be_reached_in_order_and_are_saved(self):
        self.login(self.s1)
        r = self.api(self.les.slug, 1)
        self.assertEqual(r.json()['ok'], True)
        p = LessonProgress.objects.get(user=self.s1, lesson=self.les)
        self.assertEqual((p.status, p.last_page, p.max_page), ('in_progress', 1, 1))
        self.assertIsNotNone(p.started_at)
        self.assertEqual(self.api(self.les.slug, 2).status_code, 200)
        r = self.api(self.les.slug, 4)                                  # skipping page 3
        self.assertEqual(r.status_code, 409)
        self.assertEqual(r.json()['max_page'], 2)
        p.refresh_from_db()
        self.assertEqual((p.last_page, p.max_page), (2, 2))

    def test_going_back_keeps_furthest_page(self):
        self.login(self.s1)
        for n in (1, 2, 3, 1):
            self.api(self.les.slug, n)
        p = LessonProgress.objects.get(user=self.s1, lesson=self.les)
        self.assertEqual((p.last_page, p.max_page), (1, 3))

    def test_page_outside_file_and_garbage_rejected(self):
        self.login(self.s1)
        self.assertEqual(self.api(self.les.slug, 0).status_code, 400)
        self.assertEqual(self.api(self.les.slug, 99).status_code, 400)
        self.assertEqual(self.client.post(reverse('lesson_progress', args=[self.les.slug]), 'nonsense', content_type='application/json').status_code, 400)
        self.assertEqual(self.client.get(reverse('lesson_progress', args=[self.les.slug])).status_code, 405)

    def test_progress_api_access_rules(self):
        self.login(self.s2)
        self.assertEqual(self.api(self.les.slug, 1).status_code, 404)    # wrong classroom
        self.login(self.m1)
        self.assertEqual(self.api(self.les.slug, 1).status_code, 403)    # mentors never write progress
        self.assertFalse(LessonProgress.objects.filter(lesson=self.les).exists())
        self.client.logout()
        self.assertEqual(self.api(self.les.slug, 1).status_code, 302)

    def test_resume_page_is_offered_when_reopening(self):
        self.login(self.s1)
        for n in (1, 2, 3):
            self.api(self.les.slug, n)
        r = self.client.get(reverse('lesson_detail', args=[self.les.slug]))
        self.assertEqual(r.context['resume_page'], 3)
        self.assertContains(r, 'You stopped on page 3')
        self.assertContains(r, 'data-start="3"')

    def test_complete_only_after_last_page_reached(self):
        self.login(self.s1)
        url = reverse('lesson_detail', args=[self.les.slug])
        self.api(self.les.slug, 1)
        r = self.client.post(url, {'action': 'mark_complete'}, follow=True)
        self.assertContains(r, 'read to the last page')
        self.assertFalse(LessonProgress.objects.get(user=self.s1, lesson=self.les).completed)
        for n in (2, 3):
            self.api(self.les.slug, n)
        self.client.post(url, {'action': 'mark_complete'})               # still short of the last page
        self.assertFalse(LessonProgress.objects.get(user=self.s1, lesson=self.les).completed)
        r = self.api(self.les.slug, 4)
        self.assertTrue(r.json()['can_complete'])
        self.client.post(url, {'action': 'mark_complete'})
        p = LessonProgress.objects.get(user=self.s1, lesson=self.les)
        self.assertEqual((p.status, p.completed), ('completed', True))
        self.assertIsNotNone(p.completed_at)
        self.assertTrue(UserBadge.objects.filter(user=self.s1, badge__name='First Steps').exists())

    def test_button_is_disabled_until_the_last_page(self):
        self.login(self.s1)
        html = self.client.get(reverse('lesson_detail', args=[self.les.slug])).content.decode()
        self.assertRegex(html, r'id="complete-btn"\s+disabled')

    def test_status_labels_on_the_lesson_list(self):
        self.login(self.s1)
        self.assertContains(self.client.get(reverse('lessons')), 'Not started')
        self.api(self.les.slug, 1); self.api(self.les.slug, 2)
        r = self.client.get(reverse('lessons'))
        self.assertContains(r, 'In progress')
        self.assertContains(r, 'page 2 of 4')
        for n in (3, 4):
            self.api(self.les.slug, n)
        self.client.post(reverse('lesson_detail', args=[self.les.slug]), {'action': 'mark_complete'})
        self.assertContains(self.client.get(reverse('lessons')), 'Completed')

    def test_replacing_the_file_restarts_unfinished_students_only(self):
        self.login(self.s1)
        for n in (1, 2):
            self.api(self.les.slug, n)
        done = make('student_done01', 'STUDENT'); self.join(done, self.room1)
        LessonProgress.objects.create(user=done, lesson=self.les, completed=True, max_page=4, last_page=4)
        lessonfiles.attach(self.les, upload('v2.pdf', pdf_bytes(6)), self.m1)
        p = LessonProgress.objects.get(user=self.s1, lesson=self.les)
        self.assertEqual((p.last_page, p.max_page), (1, 0))
        q = LessonProgress.objects.get(user=done, lesson=self.les)
        self.assertEqual((q.completed, q.max_page), (True, 4))

    def test_text_only_lesson_still_completes_with_the_button(self):
        les2 = lesson(self.room1, 'Text only lesson', order=9)
        self.login(self.s1)
        url = reverse('lesson_detail', args=[les2.slug])
        self.client.get(url)
        self.assertEqual(LessonProgress.objects.get(user=self.s1, lesson=les2).status, 'in_progress')
        self.client.post(url, {'action': 'mark_complete'})
        self.assertTrue(LessonProgress.objects.get(user=self.s1, lesson=les2).completed)

    def test_mentor_preview_records_nothing(self):
        self.login(self.m1)
        r = self.client.get(reverse('lesson_detail', args=[self.les.slug]))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'no progress is recorded')
        self.assertFalse(LessonProgress.objects.filter(lesson=self.les).exists())
        self.assertEqual(self.client.post(reverse('lesson_detail', args=[self.les.slug]), {'action': 'mark_complete'}).status_code, 404)


class ConversionTests(P4CBase):
    def test_no_converter_is_reported_honestly(self):
        with mock.patch.object(lessonfiles, 'find_converter', return_value=None):
            lf = lessonfiles.attach(self.les, upload('d.docx', office_zip('docx')), self.m1)
        self.assertEqual((lf.conversion_status, lf.page_count, lf.is_viewable), ('UNAVAILABLE', 0, False))
        self.assertIn('no document converter', lf.conversion_message)
        self.login(self.s1)
        r = self.client.get(reverse('lesson_detail', args=[self.les.slug]))
        self.assertContains(r, 'not ready yet')
        self.assertNotContains(r, 'id="viewer"')
        # a student is not stuck: with no readable pages the old button still completes the lesson
        self.client.post(reverse('lesson_detail', args=[self.les.slug]), {'action': 'mark_complete'})
        self.assertTrue(LessonProgress.objects.get(user=self.s1, lesson=self.les).completed)

    def test_failed_conversion_is_reported(self):
        with mock.patch.object(lessonfiles, 'find_converter', return_value='/bin/false'):
            lf = lessonfiles.attach(self.les, upload('d.docx', office_zip('docx')), self.m1)
        self.assertEqual((lf.conversion_status, lf.is_viewable), ('FAILED', False))
        self.login(self.m1)
        self.assertContains(self.client.get(reverse('mentor_lesson_edit', args=[self.les.pk])), 'Conversion failed')

    @skipUnless(soffice_available(), 'LibreOffice is not installed here')
    def test_real_word_and_powerpoint_files_become_pages(self):
        try:
            import docx, pptx
        except ImportError:
            self.skipTest('python-docx / python-pptx not installed')
        d = docx.Document(); d.add_paragraph('Hello lesson'); b = io.BytesIO(); d.save(b)
        lf = lessonfiles.attach(self.les, upload('real.docx', b.getvalue()), self.m1)
        self.assertEqual((lf.conversion_status, lf.is_viewable), ('READY', True), lf.conversion_message)
        self.assertGreaterEqual(lf.page_count, 1)
        self.login(self.s1)
        r = self.client.get(reverse('lesson_file', args=[self.les.slug]))
        self.assertTrue(b''.join(r.streaming_content).startswith(b'%PDF'))
        pr = pptx.Presentation()
        for _ in range(3):
            pr.slides.add_slide(pr.slide_layouts[5])
        b = io.BytesIO(); pr.save(b)
        lf = lessonfiles.attach(self.les, upload('real.pptx', b.getvalue()), self.m1)
        self.assertEqual((lf.conversion_status, lf.page_count), ('READY', 3), lf.conversion_message)


class QuizLinkTests(P4CBase):
    def setUp(self):
        super().setUp()
        self.qz = Quiz.objects.create(title='Lesson quiz', created_by=self.m1, classroom=self.room1, is_published=True,
                                      lesson=self.les, require_lesson_first=True)
        q = Question.objects.create(quiz=self.qz, text='Q?')
        self.right = Choice.objects.create(question=q, text='yes', is_correct=True)
        self.q = q

    def finish_lesson(self):
        self.login(self.s1)
        for n in (1, 2, 3, 4):
            self.api(self.les.slug, n)
        self.client.post(reverse('lesson_detail', args=[self.les.slug]), {'action': 'mark_complete'})

    def test_locked_until_lesson_completed_on_the_server(self):
        self.login(self.s1)
        r = self.client.get(reverse('quiz_detail', args=[self.qz.pk]))
        self.assertRedirects(r, reverse('lesson_detail', args=[self.les.slug]), fetch_redirect_response=False)
        r = self.client.post(reverse('quiz_detail', args=[self.qz.pk]), {f'question_{self.q.pk}': self.right.pk})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(QuizAttempt.objects.count(), 0)                 # forging the POST does not work

    def test_unlocked_after_completion(self):
        self.finish_lesson()
        self.assertEqual(self.client.get(reverse('quiz_detail', args=[self.qz.pk])).status_code, 200)
        self.client.post(reverse('quiz_detail', args=[self.qz.pk]), {f'question_{self.q.pk}': self.right.pk})
        self.assertEqual(QuizAttempt.objects.filter(user=self.s1).count(), 1)

    def test_lesson_page_shows_locked_then_proceed_button(self):
        self.login(self.s1)
        r = self.client.get(reverse('lesson_detail', args=[self.les.slug]))
        self.assertContains(r, 'Locked')
        self.assertNotContains(r, 'Proceed to Quiz')
        self.finish_lesson()
        r = self.client.get(reverse('lesson_detail', args=[self.les.slug]))
        self.assertContains(r, 'Proceed to Quiz')
        self.assertContains(r, reverse('quiz_detail', args=[self.qz.pk]))

    def test_quiz_without_the_rule_is_always_open_but_still_linked(self):
        Quiz.objects.filter(pk=self.qz.pk).update(require_lesson_first=False)
        self.login(self.s1)
        self.assertEqual(self.client.get(reverse('quiz_detail', args=[self.qz.pk])).status_code, 200)
        self.assertContains(self.client.get(reverse('lesson_detail', args=[self.les.slug])), 'Proceed to Quiz')

    def test_quiz_list_marks_locked_quizzes(self):
        self.login(self.s1)
        self.assertContains(self.client.get(reverse('quizzes')), 'Finish the lesson first')
        self.finish_lesson()
        self.assertNotContains(self.client.get(reverse('quizzes')), 'Finish the lesson first')

    def test_mentor_and_admin_are_not_blocked(self):
        for u in (self.m1, self.admin):
            self.login(u)
            self.assertEqual(self.client.get(reverse('quiz_detail', args=[self.qz.pk])).status_code, 200)

    def test_other_students_quiz_progress_is_not_shared(self):
        self.finish_lesson()
        s3 = make('student_ccc3', 'STUDENT'); self.join(s3, self.room1)
        self.login(s3)
        self.assertEqual(self.client.get(reverse('quiz_detail', args=[self.qz.pk])).status_code, 302)

    def test_mentor_form_requires_a_lesson_for_the_rule(self):
        self.login(self.m1)
        data = {'classroom': self.room1.pk, 'title': 'T', 'description': '', 'instructions': '', 'require_lesson_first': 'on'}
        r = self.client.post(reverse('mentor_quiz_create'), data)
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'Choose a linked lesson first')
        data['lesson'] = self.les.pk
        r = self.client.post(reverse('mentor_quiz_create'), data)
        self.assertEqual(r.status_code, 302)
        self.assertTrue(Quiz.objects.get(title='T').require_lesson_first)


class MigrationCompatTests(P4CBase):
    def test_old_lesson_progress_rows_still_work(self):
        les2 = lesson(self.room1, 'Old lesson', order=8)
        p = LessonProgress.objects.create(user=self.s1, lesson=les2, completed=True)
        self.assertEqual((p.status, p.last_page, p.max_page), ('completed', 1, 0))
        self.login(self.s1)
        self.assertEqual(self.client.get(reverse('lesson_detail', args=[les2.slug])).status_code, 200)
        self.assertContains(self.client.get(reverse('lessons')), 'Completed')
